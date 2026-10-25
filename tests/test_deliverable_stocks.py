"""📦 Deliverable Stocks (deliverable_stocks.py): the date helpers, the chart series, the old workbook's 'Deliverable Stox' tab and a
save/load round trip on a throw-away SQLite file (never the real archive).

    python tests/test_deliverable_stocks.py
"""
import datetime as dt
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
os.environ["USE_SNOWFLAKE"] = ""          # db.py reads these on every call: keep the test off Snowflake / Postgres
os.environ["DATABASE_URL"] = ""

import db                                  # noqa: E402
import deliverable_stocks as ds            # noqa: E402

FAILS = []
D = dt.date


def check(name, cond, detail=""):
    print("  [%s] %s%s" % ("PASS" if cond else "FAIL", name, ("  -- " + str(detail)) if (detail and not cond) else ""))
    if not cond:
        FAILS.append(name)


print("dates")
check("last_friday of a Friday is itself", ds.last_friday(D(2026, 10, 23)) == D(2026, 10, 23))
check("last_friday of a Thursday is the week before", ds.last_friday(D(2026, 10, 22)) == D(2026, 10, 16))
check("Tue Oct 27 -> the Oct 23 report is out", ds.newest_report_friday(D(2026, 10, 27)) == D(2026, 10, 23))
check("Mon Oct 26 -> still Oct 16", ds.newest_report_friday(D(2026, 10, 26)) == D(2026, 10, 16))
check("Sat Oct 24 -> Oct 16", ds.newest_report_friday(D(2026, 10, 24)) == D(2026, 10, 16))

# the archive as the workbook leaves it: corn to Jun 19, soybeans to Jun 12
data = {"Corn": {D(2026, 6, 12): 1.0, D(2026, 6, 19): 1.0}, "Soybeans": {D(2026, 6, 12): 1.0}}
miss = ds.missing_fridays(data, D(2026, 10, 16))
check("missing starts at the oldest gap (soybeans Jun 19)", miss[:2] == [D(2026, 6, 19), D(2026, 6, 26)], miss[:3])
check("missing runs to the newest report, 18 Fridays", len(miss) == 18 and miss[-1] == D(2026, 10, 16), (len(miss), miss[-1:]))
check("a week typed for both commodities is not missing",
      D(2026, 6, 26) not in ds.missing_fridays({c: {**s, D(2026, 6, 26): 2.0} for c, s in data.items()}, D(2026, 10, 16)))
check("missing from a Wednesday starts at the next Friday",
      ds.missing_fridays(data, D(2026, 7, 10), since=D(2026, 7, 1)) == [D(2026, 7, 3), D(2026, 7, 10)])
check("nothing entered -> nothing listed", ds.missing_fridays({}, D(2026, 10, 16)) == [])

print("chart series")
series = {}
for y in range(2019, 2027):
    f = ds._first_friday(y)
    for k in range(52 if y < 2026 else 24):
        series[f + dt.timedelta(days=7 * k)] = 1000.0 * (y - 2018) + k
check("chart years = six back + current", ds.chart_years(series, 2026) == [2020, 2021, 2022, 2023, 2024, 2025, 2026])
check("average years = the five before the current", ds.average_years(series, 2026) == [2021, 2022, 2023, 2024, 2025])
check("a year with no data is left out", ds.chart_years({D(2024, 1, 5): 1.0, D(2026, 1, 2): 1.0}, 2026) == [2024, 2026])
pts = ds.season_points(series, [2024])
check("season points sit on the REF_YEAR axis", all(p["x"].year == ds.REF_YEAR for p in pts) and pts[0]["date"] == D(2024, 1, 5))
check("Feb 29 maps onto the leap REF_YEAR", ds.season_points({D(2024, 2, 29): 5.0}, [2024])[0]["x"] == D(2000, 2, 29))
avg = ds.average_points(series, [2021, 2022, 2023, 2024, 2025])
first = [p for p in avg if p["x"].month == 1][0]
check("average of week 1 = mean of the five years", abs(first["stocks"] - sum(1000.0 * (y - 2018) for y in range(2021, 2026)) / 5) < 1e-9,
      first)
check("a week needs three years", ds.average_points({D(2024, 1, 5): 1.0, D(2025, 1, 3): 3.0}, [2024, 2025]) == [])

print("workbook")
import openpyxl                            # noqa: E402

tmpdir = tempfile.mkdtemp()
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Deliverable Stox"
years = list(range(2026, 2019, -1))        # header 2026 .. 2020, then the average column
r = 1
for name, base in (("Soybeans", 10000), ("Corn", 50000)):
    ws.cell(r, 1, name)
    ws.cell(r + 1, 2, "Week")
    for c, y in enumerate(years, start=3):
        ws.cell(r + 1, c, y)
    ws.cell(r + 1, 3 + len(years), "Avg")
    first = ds._first_friday(2026)
    for i, k in enumerate(range(51, -1, -1)):          # Dec at the top, Jan at the bottom, like the workbook
        row = r + 2 + i
        ws.cell(row, 2, dt.datetime.combine(first + dt.timedelta(days=7 * k), dt.time()))
        for c, y in enumerate(years, start=3):
            if y == 2026 and k > 23:                   # the current year stops in mid-June
                continue
            ws.cell(row, c, base + 100 * (2026 - y) + k)
        ws.cell(row, 3 + len(years), 999999)           # the average column: never read
    r += 2 + 52 + 3
path = os.path.join(tmpdir, "stox.xlsx")
wb.save(path)
rows = ds.parse_workbook(path)
got = {(a, c): v for a, c, v, s in rows}
check("every row is tagged 'workbook'", {s for *_x, s in rows} == {"workbook"})
check("2026 stops after 24 weeks", sum(1 for (a, c) in got if a.year == 2026 and c == "Corn") == 24)
check("past years have 52 weeks", sum(1 for (a, c) in got if a.year == 2021 and c == "Soybeans") == 52)
check("2026 week 0 = Fri Jan 2", got[(D(2026, 1, 2), "Corn")] == 50000.0)
check("row k of a past year is dated that year's k-th Friday (2020 row 0 = Jan 3)", got[(D(2020, 1, 3), "Soybeans")] == 10600.0)
check("2021 row 51 = Dec 24 2021", got[(D(2021, 12, 24), "Corn")] == 50000.0 + 500 + 51)
check("the average column is not read", 999999.0 not in got.values())
check("both blocks are read", {c for (_a, c) in got} == {"Corn", "Soybeans"})

print("save / load (temp SQLite)")
db.LOCAL_SQLITE = os.path.join(tmpdir, "ds.db")
check("no table yet -> {}", ds.load_all() == {})
n = ds.save([(D(2026, 6, 19), "Corn", 17852.0, "workbook"), (D(2026, 6, 19), "Soybeans", None, "workbook"),
             (D(2026, 6, 19), "Wheat", 5.0, "workbook"), ("2026-06-26", "Corn", 18000, "manual")])
check("None and unknown commodities are skipped", n == 2, n)
check("loads back as dates and floats", ds.load_all() == {"Corn": {D(2026, 6, 19): 17852.0, D(2026, 6, 26): 18000.0}})
ds.save([(D(2026, 6, 19), "Corn", 17900.0, "manual")])
check("typing a week again replaces it", ds.load_all()["Corn"][D(2026, 6, 19)] == 17900.0)
conn, _ = db._connect()
src = dict(conn.execute(f"SELECT as_of, source FROM {ds.TABLE} WHERE commodity = 'Corn'").fetchall())
conn.close()
check("the source follows the latest save", src == {"2026-06-19": "manual", "2026-06-26": "manual"}, src)

print()
if FAILS:
    print(f"{len(FAILS)} FAILED: {FAILS}")
    sys.exit(1)
print("all passed")
